"""Workbook of one round's markets: Summary, New Qualified, All New Leads, LinkedIn Leads.

    python3 scripts/export_run.py OUT.xlsx "City, ST" ...
"""
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font

from common import COLUMNS, RUN_DIR
from export_xlsx import WIDTHS, read, sheet


def main():
    out, markets = sys.argv[1], sys.argv[2:]
    rows = [r for r in read("diamonds_all.csv") if r["Market"] in markets]
    qual = [r for r in read("diamonds_qualified.csv") if r["Market"] in markets]
    li = [r for r in rows if r["LinkedIn Profile"]]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Market", "Businesses", "Qualified (2 of 3)", "With LinkedIn", "With email"])
    for c in ws[1]:
        c.font = Font(bold=True)
    qc = Counter(r["Market"] for r in qual)
    for m in markets:
        rs = [r for r in rows if r["Market"] == m]
        ws.append([m, len(rs), qc[m], sum(bool(r["LinkedIn Profile"]) for r in rs),
                   sum(bool(r["Email"]) for r in rs)])
    ws.append(["TOTAL", len(rows), len(qual), len(li), sum(bool(r["Email"]) for r in rows)])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append([])
    ws.append(["Note: suburb directory searches mostly return NYC / Philadelphia / Chicago firms; those are "
               "deduped into their real market and are not counted here. Every row has a real Source URL."])
    ws.column_dimensions["A"].width = 26
    for col in "BCDE":
        ws.column_dimensions[col].width = 18
    sheet(wb, "New Qualified", COLUMNS, qual, WIDTHS)
    sheet(wb, "All New Leads", COLUMNS, rows, WIDTHS)
    sheet(wb, "LinkedIn Leads", COLUMNS, li, WIDTHS)
    wb.save(RUN_DIR / out)
    print(f"wrote {out}: all={len(rows)} qualified={len(qual)} linkedin={len(li)}")
    for m in markets:
        rs = [r for r in rows if r["Market"] == m]
        print(f"  {m:24}{len(rs):4}{qc[m]:4}{sum(bool(r['LinkedIn Profile']) for r in rs):4}")


if __name__ == "__main__":
    main()
