"""Compile the built CSVs into leadrun_diamonds/diamond_leads.xlsx (one workbook, several sheets).

Sheets: Summary, Qualified (2 of 3), All Leads, Dropped (out of scope).
Reads only the builder's outputs, so run build.py first. Never touches the production
usaENRICHEDleads.xlsx.
"""
import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from common import COLUMNS, RUN_DIR

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WIDTHS = {"Business": 34, "Category (as listed)": 24, "Phone": 16, "Email": 30, "LinkedIn Profile": 42,
          "Contact Person": 22, "Role": 14, "Website": 32, "Address": 38, "Real City": 16, "Real State": 9,
          "Market": 18, "Lead Origin": 16, "Source": 48, "Signals": 40, "Flags": 24,
          "Building / Suite": 30, "Trade Memberships": 22, "Mobile Candidates": 20, "Socials": 36}


def read(name):
    with open(RUN_DIR / name, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append([r.get(h, "") if isinstance(r, dict) else r[i] for i, h in enumerate(header)])
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 14)
    ws.freeze_panes = "B2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    # make URL cells clickable
    for col in ("Source", "Website", "LinkedIn Profile", "Email Source", "LinkedIn Source"):
        if col in header:
            j = header.index(col) + 1
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                c = row[0]
                if isinstance(c.value, str) and c.value.startswith("http"):
                    c.hyperlink, c.font = c.value, Font(color="0563C1", underline="single")
    return ws


def main():
    all_rows, qual = read("diamonds_all.csv"), read("diamonds_qualified.csv")
    dropped = read("dropped_out_of_scope.csv")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    n = len(all_rows)
    lines = [
        ("Diamond trade lead run", ""),
        ("Unique businesses researched (in scope)", n),
        ("Qualified (2 of 3: email / personal LinkedIn / direct line)", len(qual)),
        ("Dropped out of scope", len(dropped)),
        ("With phone", sum(bool(r["Phone"]) for r in all_rows)),
        ("With real email", sum(bool(r["Email"]) for r in all_rows)),
        ("With personal LinkedIn", sum(bool(r["LinkedIn Profile"]) for r in all_rows)),
        ("With named contact", sum(bool(r["Contact Person"]) for r in all_rows)),
        ("With trade membership", sum(bool(r["Trade Memberships"]) for r in all_rows)),
        ("", ""), ("Per market", "Businesses"),
    ]
    lines += Counter(r["Market"] for r in all_rows).most_common()
    lines += [("", ""), ("Per lead origin", "Businesses")]
    lines += Counter(r["Lead Origin"] for r in all_rows).most_common()
    lines += [("", ""), ("Dropped by reason", "Count")]
    lines += Counter(r["Reason"] for r in dropped).most_common()
    lines += [("", ""), ("Notes", "See REPORT.md: site crawling was blocked by the sandbox network policy, "
                                 "so email / direct-line fill is near zero. Every row has a real Source URL.")]
    for a, b in lines:
        ws.append([a, b])
    for r in (1, 11):
        ws.cell(r, 1).font = Font(bold=True, size=13 if r == 1 else 11)
    for row in ws.iter_rows():
        if row[0].value in ("Per lead origin", "Dropped by reason", "Notes"):
            row[0].font = Font(bold=True)
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 60

    sheet(wb, "Qualified", COLUMNS, qual, WIDTHS)
    sheet(wb, "All Leads", COLUMNS, all_rows, WIDTHS)
    sheet(wb, "Dropped", ["Business", "Market", "Reason", "Stage", "Source", "Why"], dropped,
          {"Business": 34, "Market": 18, "Reason": 16, "Source": 48, "Why": 70})
    out = RUN_DIR / "diamond_leads.xlsx"
    wb.save(out)
    print(f"wrote {out} — all={n} qualified={len(qual)} dropped={len(dropped)}")


if __name__ == "__main__":
    main()
