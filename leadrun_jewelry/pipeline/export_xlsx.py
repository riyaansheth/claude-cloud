"""jewelry_qualified.xlsx — the qualified leads with every column, plus summary
sheets. Built from the CSVs the builder just wrote (no other data source)."""
import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .util import now_iso

WIDTHS = {"Business": 32, "Segment": 22, "Category (as listed)": 30, "Phone": 16,
          "Email": 32, "LinkedIn Profile": 34, "Contact Person": 20, "Role": 14,
          "Website": 30, "Address": 40, "Real City": 16, "Real State": 9, "Market": 16,
          "Lead Origin": 14, "Rating": 8, "Reviews": 9, "Source": 40, "Signals": 40,
          "Match Type": 11, "Email Source": 40, "LinkedIn Source": 30, "Found At": 26,
          "Also Listed As": 30, "Other Markets": 18, "Merged Listings": 9, "Flags": 20,
          "Mobile Candidates": 30, "Socials": 40}
LINK_COLS = {"Website", "Source", "Email Source", "LinkedIn Profile", "LinkedIn Source"}
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def _read(path):
    if not path.exists():
        return []
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def _table(ws, rows):
    for r in rows:
        ws.append(r)
    for i, col in enumerate(ws[1], 1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(45, len(str(col.value)) + 4))


def write():
    q = _read(config.QUALIFIED_CSV)
    allrows = _read(config.ALL_CSV)
    wb = Workbook()

    ws = wb.active
    ws.title = "Qualified Leads"
    _header(ws, config.COLUMNS)
    for r in q:
        vals = []
        for c in config.COLUMNS:
            v = r.get(c, "")
            if c in ("Rating",) and v:
                try:
                    v = float(v)
                except ValueError:
                    pass
            elif c in ("Reviews", "Merged Listings") and v:
                try:
                    v = int(v)
                except ValueError:
                    pass
            vals.append(v)
        ws.append(vals)
    for i, c in enumerate(config.COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 18)
        if c in LINK_COLS:
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=i)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = Font(color="0563C1", underline="single")
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("By Market")
    _header(s, ["Market", "Researched (all.csv)", "Qualified", "Rate", "With email",
                "With owner name", "With LinkedIn"])
    mk = [config.market_label(c, st) for c, st, _ in config.MARKETS]
    ca, cq = Counter(r["Market"] for r in allrows), Counter(r["Market"] for r in q)
    rows = []
    for m in mk:
        if ca.get(m) or cq.get(m):
            sub = [r for r in q if r["Market"] == m]
            rows.append([m, ca.get(m, 0), cq.get(m, 0),
                         f"{100 * cq.get(m, 0) / ca[m]:.1f}%" if ca.get(m) else "",
                         sum(bool(r["Email"]) for r in sub), sum(bool(r["Contact Person"]) for r in sub),
                         sum(bool(r["LinkedIn Profile"]) for r in sub)])
    rows.append(["TOTAL", len(allrows), len(q), f"{100 * len(q) / len(allrows):.1f}%" if allrows else "",
                 sum(bool(r["Email"]) for r in q), sum(bool(r["Contact Person"]) for r in q),
                 sum(bool(r["LinkedIn Profile"]) for r in q)])
    _table(s, rows)

    s = wb.create_sheet("By Segment & Source")
    _header(s, ["Segment", "Qualified"])
    _table(s, sorted(Counter(r["Segment"] for r in q).items(), key=lambda x: -x[1]))
    s.append([])
    s.append(["Lead Origin", "Qualified"])
    for k, v in sorted(Counter(r["Lead Origin"] for r in q).items(), key=lambda x: -x[1]):
        s.append([k, v])

    s = wb.create_sheet("Read Me")
    notes = [
        ["Generated", now_iso()],
        ["Qualified leads", len(q)],
        ["Businesses researched", len(allrows)],
        ["Qualification bar", "2 of 3: real email · personal LinkedIn (/in/, named human) · "
                              "direct/mobile number from a tel: link that is not any directory-listed number"],
        ["Signals", "'N of 3' + tags: role-email (info@/sales@…), direct-line, gia, ags, ja, bench, "
                    "custom-design, since-YYYY"],
        ["Mobile Candidates", "Every US tel: number on the business's own site/socials except listed lines"],
        ["Scope", "Independent jewelry retail only. Kill-list (pawn, gold buyers, coins, beads/craft, "
                  "costume/fashion, piercing, kiosks, simulants, chains, thrift, appraisers-only, "
                  "diamond trade) applied at discovery, after the site crawl and in the builder."],
        ["Provenance", "Every row has a real Source URL; every email has the exact page it came from. "
                       "Nothing is filled from memory."],
    ]
    for n in notes:
        s.append(n)
    s.column_dimensions["A"].width = 22
    s.column_dimensions["B"].width = 120
    for row in s.iter_rows():
        row[0].font = Font(bold=True)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(config.QUALIFIED_XLSX)
    return config.QUALIFIED_XLSX
